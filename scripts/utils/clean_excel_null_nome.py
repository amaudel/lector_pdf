import pandas as pd
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

root_dir = Path("c:/Users/User/Documents/LECTOR PDF")
files = [root_dir / "reporte_acumulado.xlsx", root_dir / "reporte_acumulado_1.xlsx"]

def clean_excel_file(file_path):
    if not file_path.exists():
        return
        
    print(f"Saneando archivo Excel de registros null/nome/placeholders: {file_path.name}")
    try:
        # 1. Leer con pandas para filtrar
        df = pd.read_excel(file_path, sheet_name="Master", dtype=str)
        df = df.fillna("")
        
        placeholders = [
            "NONE", "NULL", "NOME", "UNKNOWN", "DESCONOCIDA", "DESCONOCIDO", "SIN DATO", "SIN_DATO", "NO DISPONIBLE"
        ]
        
        # 2. Filtrar filas de entidades basura
        def is_valid_row(row):
            rz = str(row.get("RAZÓN SOCIAL", "")).strip().upper()
            
            # Exclusión por longitud
            if len(rz) < 3:
                return False
                
            # Exclusión de placeholders
            for p in placeholders:
                if rz == p or rz.startswith(p + " "):
                    return False
                    
            return True
            
        initial_len = len(df)
        df_filtered = df[df.apply(is_valid_row, axis=1)].copy()
        final_len = len(df_filtered)
        
        print(f"  Filas iniciales: {initial_len} | Filas conservadas: {final_len} | Eliminadas: {initial_len - final_len}")
        
        if final_len == 0:
            print("  El archivo quedaría vacío, omitiendo guardado.")
            return
            
        # Re-secuenciar el Nro
        df_filtered["Nro"] = [str(i) for i in range(1, final_len + 1)]
        
        # 3. Escribir con openpyxl aplicando los estilos correctos
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df_filtered.to_excel(writer, index=False, sheet_name="Master")
            wb = writer.book
            ws = writer.sheets["Master"]
            
            # Estilos
            col_colors = {
                1: 'FCE4D6', # Nro
                2: 'D9E1F2', # MEMORANDO INICIAL
                3: 'D9E1F2', # MEMORANDO DP
                4: 'E7E6E6', # SOLICITANTE
                5: 'F8CBAD', # RAZÓN SOCIAL
                6: 'F8CBAD', # RUC
                7: 'BDD7EE', # SUCURSAL
                8: 'BDD7EE', # RESOLUCIÓN
                9: '5B9BD5', # PROVINCIA
                10: 'C6E0B4', # ACCIONES
                11: 'C65911'  # OBSERVACIONES
            }
            
            header_font = Font(name='Calibri', bold=True, color='000000', size=11)
            header_font_white = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
            header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell_align = Alignment(vertical='center', wrap_text=False)
            cell_align_left = Alignment(horizontal='left', vertical='center', wrap_text=False)
            cell_align_center = Alignment(horizontal='center', vertical='center', wrap_text=False)
            
            thin_border = Border(
                left=Side(style='thin', color='7F7F7F'),
                right=Side(style='thin', color='7F7F7F'),
                top=Side(style='thin', color='7F7F7F'),
                bottom=Side(style='thin', color='7F7F7F')
            )
            
            # Formatear Cabecera
            ws.row_dimensions[1].height = 28
            for col_idx in range(1, 12):
                cell = ws.cell(row=1, column=col_idx)
                if col_idx in [9, 11]:
                    cell.font = header_font_white
                else:
                    cell.font = header_font
                color = col_colors.get(col_idx, 'D9D9D9')
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
                cell.alignment = header_align
                cell.border = thin_border
                
            # Formatear Celdas de Datos
            yellow_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
            
            for row_idx in range(2, len(df_filtered) + 2):
                ws.row_dimensions[row_idx].height = 20
                for col_idx in range(1, 12):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(name='Calibri', size=10)
                    cell.border = thin_border
                    
                    # Rellenos amarillos para E, F, I
                    if col_idx in [5, 6, 9]:
                        cell.fill = yellow_fill
                        
                    # Alineaciones y formatos
                    if col_idx in [1, 6, 7, 9]:  # Nro, RUC, Sucursal, Provincia
                        cell.alignment = cell_align_center
                    else:
                        cell.alignment = cell_align_left
                        
                    if col_idx in [6, 7]:  # RUC and SUCURSAL strictly text
                        cell.number_format = '@'
                        
            ws.freeze_panes = 'A2'
            
            # Anchos
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    val = str(cell.value or '')
                    if len(val) > max_len:
                        max_len = len(val)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
                
        print(f"  [OK] Guardado archivo saneado: {file_path.name}")
        
    except Exception as e:
        print(f"  Error saneando {file_path.name}: {e}")

for f in files:
    clean_excel_file(f)
